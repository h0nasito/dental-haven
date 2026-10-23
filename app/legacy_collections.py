"""MyMedsPH collection reports (Excel, one per year) -> payments on the converted MyMedsPH bills.

The MyMedsPH data export has bills but no payments. Its Reports > Collections export does: a workbook with a
"Collection Report" sheet (one row per payment: Date, Patient, Clinic/Lab, Reference #, Remarks, Amount) and one sheet
per dentist with the same rows plus "Procedures". This module reads those workbooks and:

- matches each payment to a patient by name (the report has no patient ID); names that match no patient, or more than
  one, are listed and skipped;
- puts the patient's MyMedsPH bills into real bills (all of them, unpaid; "keep as history" statuses stay out);
- applies each payment to the patient's bills: first a bill with a line for the procedure paid for, else the oldest
  bill with a balance; a payment bigger than what is owed is split over several bills, and what is left after every
  bill is paid becomes account credit;
- sets the dentist on bill lines when the dentist's sheet shows the procedure (for commission);
- keeps a fingerprint of each row, so uploading the same or an overlapping report again adds nothing twice.
Payment type: "cash", "gcash", bank names (BDO, BPI, ...) and cards are mapped; anything else is "Other" with the text kept.
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import defaultdict
from datetime import datetime

from . import legacy_billing, settings
from .billing import banks
from .util import now_str

REPORT_SHEET = "Collection Report"
HEAD = ["date", "patient", "clinic/lab", "reference #", "remarks", "amount"]


class CollectionFileError(Exception):
    pass


def norm(s: str) -> str:
    s = (s or "").lower().replace("ñ", "n").replace("Ñ", "n")
    s = re.sub(r"\b(dr|dra|dmd)\b\.?", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def _num(v) -> int | None:
    try:
        return round(float(str(v).replace(",", "").strip()) * 100)
    except (TypeError, ValueError):
        return None


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    for fmt in ("%b %d, %Y", "%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except (TypeError, ValueError):
            continue
    return None


def _cells(r) -> list[str]:
    return ["" if v is None else str(v).strip() for v in r]


def read_reports(files: list[tuple[str, bytes]]) -> dict:
    """Read MyMedsPH collection report workbooks -> {"rows": [...], "files": [...], "errors": [...]}."""
    import openpyxl
    rows, info, errors = [], [], []
    seen_count: dict[str, int] = defaultdict(int)
    for name, data in files:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable file
            raise CollectionFileError(f"{name}: this isn't an Excel file MyMedsPH made ({exc.__class__.__name__}).") from exc
        if REPORT_SHEET not in wb.sheetnames:
            raise CollectionFileError(f"{name}: no '{REPORT_SHEET}' sheet. Export it from MyMedsPH Reports > Collections.")
        # dentist sheets: same rows plus Procedures; the sheet name is the dentist
        by_dentist: dict[tuple, list] = defaultdict(list)
        for ws in wb.worksheets:
            if ws.title in (REPORT_SHEET, "Collection Summary"):
                continue
            it = ws.iter_rows(values_only=True)
            head = [c.lower() for c in _cells(next(it, []))]
            if head[:6] != HEAD:
                continue
            for r in it:
                c = _cells(r) + [""] * 7
                by_dentist[tuple(c[:6])].append((ws.title.strip(), c[6]))
        it = wb[REPORT_SHEET].iter_rows(values_only=True)
        head = [c.lower() for c in _cells(next(it, []))]
        if head[:6] != HEAD:
            raise CollectionFileError(f"{name}: the '{REPORT_SHEET}' sheet doesn't have the usual columns ({', '.join(HEAD)}).")
        n = total = 0
        for line, r in enumerate(it, start=2):
            c = (_cells(r) + [""] * 6)[:6]
            if not any(c):
                continue
            d, amount = _date(c[0]), _num(c[5])
            if not d or amount is None or not c[1]:
                errors.append(f"{name} line {line}: missing date, patient or amount - skipped")
                continue
            if amount <= 0:
                continue
            base = "|".join([d.isoformat(), norm(c[1]), norm(c[2]), c[3].lower(), c[4].lower(), str(amount)])
            seen_count[base] += 1
            key = "mmc:" + hashlib.sha1(f"{base}#{seen_count[base]}".encode()).hexdigest()[:24]
            pool = by_dentist.get(tuple(c[:6])) or []
            dentist, procs = pool.pop(0) if pool else ("", "")
            rows.append({"date": d.isoformat(), "patient": c[1], "clinic": c[2], "ref": c[3], "remarks": c[4].lstrip("-").strip(),
                         "amount": amount, "dentist": dentist, "procedures": procs, "key": key})
            n += 1
            total += amount
        info.append({"file": name, "rows": n, "total": total})
    return {"rows": rows, "files": info, "errors": errors}


# ---------------------------------------------------------------------------
# matching
# ---------------------------------------------------------------------------

def patient_index(conn) -> dict[str, list[int]]:
    """Normalized name -> patient ids. A name is stored as 'first last' and 'first middle last'."""
    idx: dict[str, set] = defaultdict(set)
    for p in conn.all("SELECT id, first_name, middle_name, last_name FROM patients WHERE active = 1 OR active IS NULL"):
        first, mid, last = norm(p["first_name"]), norm(p["middle_name"] or ""), norm(p["last_name"])
        for k in {f"{first} {last}", f"{first} {mid} {last}" if mid else "", f"{last} {first}"}:
            if k.strip():
                idx[" ".join(k.split())].add(p["id"])
    return {k: sorted(v) for k, v in idx.items()}


def _branch_for(conn):
    branches = conn.all("SELECT id, name, slug FROM branches")
    cache: dict[str, int | None] = {}

    def find(clinic: str):
        k = norm(clinic)
        if k not in cache:
            hit = None
            for b in branches:
                if norm(b["name"]) and (norm(b["name"]) in k or k.endswith(norm(b["slug"] or "#"))):
                    hit = b["id"]
                    break
            cache[k] = hit
        return cache[k]
    return find


def _method(ref: str, conn) -> tuple[str, str]:
    r = norm(ref)
    if not r:
        return "other", ""
    if r in ("cash", "cashbdo") or r.startswith("cash "):
        return "cash", ""
    if "gcash" in r or r in ("gash", "g cash"):
        return "gcash", "GCash"
    if "maya" in r or "paymaya" in r:
        return "maya", ""
    if any(w in r for w in ("card", "mastercard", "visa", "credit", "debit")):
        return "card", ""
    for b in banks(conn) + ["Metrobank", "Unionbank", "Citi", "Landbank", "PNB", "China Bank", "EastWest"]:
        if norm(b) and (norm(b) in r or r in norm(b)):
            return "bank_transfer", b
    if r in ("security",):
        return "bank_transfer", "Security Bank"
    if r in ("union bank",):
        return "bank_transfer", "Unionbank"
    if r in ("bank", "bank transfer", "online transfer", "transfer"):
        return "bank_transfer", ""
    return "other", ref[:60]


def preview(conn, data: dict) -> dict:
    idx = patient_index(conn)
    have = {r["legacy_key"] for r in conn.all("SELECT legacy_key FROM payments WHERE legacy_key IS NOT NULL")} | \
           {r["legacy_key"] for r in conn.all("SELECT legacy_key FROM patient_credits WHERE legacy_key IS NOT NULL")}
    s = {"rows": len(data["rows"]), "total": sum(r["amount"] for r in data["rows"]), "matched": 0, "matched_total": 0, "already": 0,
         "unmatched": 0, "ambiguous": 0, "unmatched_total": 0, "patients": set(), "unmatched_names": defaultdict(lambda: [0, 0]),
         "dentists": set(), "files": data["files"], "errors": data["errors"]}
    for r in data["rows"]:
        if r["key"] in have:
            s["already"] += 1
            continue
        ids = idx.get(norm(r["patient"]), [])
        if len(ids) == 1:
            s["matched"] += 1
            s["matched_total"] += r["amount"]
            s["patients"].add(ids[0])
        else:
            s["ambiguous" if ids else "unmatched"] += 1
            s["unmatched_total"] += r["amount"]
            u = s["unmatched_names"][(r["patient"], "several patients" if ids else "no patient")]
            u[0] += 1
            u[1] += r["amount"]
        if r["dentist"]:
            s["dentists"].add(r["dentist"])
    users = {norm(u["name"]): u["id"] for u in conn.all("SELECT id, name FROM users WHERE role = 'dentist'")}
    s["dentists_unmatched"] = sorted(d for d in s["dentists"] if norm(d) not in users)
    s["patients"] = len(s["patients"])
    s["unmatched_names"] = sorted(((k[0], k[1], v[0], v[1]) for k, v in s["unmatched_names"].items()), key=lambda x: -x[3])
    s["dentists"] = len(s["dentists"])
    return s


# ---------------------------------------------------------------------------
# applying
# ---------------------------------------------------------------------------

def apply(conn, data: dict, user_id: int) -> dict:
    """Convert the bills, add the payments, set dentists on lines. Call inside a transaction."""
    st = {"payments": 0, "paid_total": 0, "credit": 0, "credit_total": 0, "skipped": 0, "already": 0, "lines_dentist": 0, "bills": 0}
    # 1. every MyMedsPH bill line not set aside as history becomes a bill (unpaid; payments come next)
    smap = legacy_billing.saved_map(conn)
    codes = [r["c"] for r in conn.all("SELECT DISTINCT status_code AS c FROM legacy_bills WHERE invoice_id IS NULL")]
    mapping = {c: ("history" if smap.get(c) == "history" else "unpaid") for c in codes}
    st["bills"] = legacy_billing.convert(conn, mapping, settings.get(legacy_billing.BRANCH_KEY, conn), user_id, with_payments=False)["bills"]
    settings.put(legacy_billing.MAP_KEY, {**smap, **{c: v for c, v in mapping.items()}}, user_id, conn)
    settings.put(legacy_billing.COLLECTIONS_KEY, True, user_id, conn)
    # 2. what each patient still owes, bill by bill (only MyMedsPH bills)
    idx = patient_index(conn)
    have = {r["legacy_key"] for r in conn.all("SELECT legacy_key FROM payments WHERE legacy_key IS NOT NULL")} | \
           {r["legacy_key"] for r in conn.all("SELECT legacy_key FROM patient_credits WHERE legacy_key IS NOT NULL")}
    bills: dict[int, list[dict]] = defaultdict(list)
    for b in conn.all("SELECT i.id, i.patient_id, i.branch_id, substr(i.issued_at, 1, 10) AS day, i.total_cents, "
                      "COALESCE((SELECT SUM(CASE WHEN p.kind = 'payment' THEN p.amount_cents ELSE -p.amount_cents END) FROM payments p "
                      "WHERE p.invoice_id = i.id AND p.status = 'valid'), 0) AS paid FROM invoices i WHERE i.legacy = 1 AND i.status = 'issued' "
                      "ORDER BY i.patient_id, i.issued_at, i.id"):
        bills[b["patient_id"]].append({**b, "due": b["total_cents"] - b["paid"], "lines": []})
    line_rows = conn.all("SELECT ii.id, ii.invoice_id, lower(ii.description) AS d, ii.dentist_id FROM invoice_items ii "
                         "JOIN invoices i ON i.id = ii.invoice_id WHERE i.legacy = 1")
    by_inv: dict[int, list] = defaultdict(list)
    for ln in line_rows:
        by_inv[ln["invoice_id"]].append(dict(ln))
    for blist in bills.values():
        for b in blist:
            b["lines"] = by_inv.get(b["id"], [])
    users = {norm(u["name"]): u["id"] for u in conn.all("SELECT id, name FROM users WHERE role = 'dentist'")}
    branch_of = _branch_for(conn)
    ts = now_str()

    def procs(text: str) -> list[str]:
        return [norm(re.sub(r"\s*\(#.*$", "", p)) for p in re.split(r"\),\s*", text or "") if p.strip()]

    def same(line_desc: str, wanted: list[str]) -> bool:
        d = norm(line_desc)
        return bool(d) and any(w and (d.startswith(w) or w.startswith(d)) for w in wanted)

    for r in sorted(data["rows"], key=lambda x: (x["date"], x["key"])):
        if r["key"] in have:
            st["already"] += 1
            continue
        ids = idx.get(norm(r["patient"]), [])
        if len(ids) != 1:
            st["skipped"] += 1
            continue
        pid, left = ids[0], r["amount"]
        plist = bills.get(pid, [])
        method, account = _method(r["ref"], conn)
        note = "MyMedsPH payment" + (f" · {r['remarks']}" if r["remarks"] else "") + (f" · {r['dentist']}" if r["dentist"] else "") \
            + (f" · {r['procedures']}" if r["procedures"] else "")
        wanted = procs(r["procedures"])
        dentist_id = users.get(norm(r["dentist"])) if r["dentist"] else None
        # bills to pay: one with a line for the procedure first, then oldest with a balance
        order = [b for b in plist if b["due"] > 0 and any(same(ln["d"], wanted) for ln in b["lines"])]
        order += [b for b in plist if b["due"] > 0 and b not in order]
        part = 0
        for b in order:
            if left <= 0:
                break
            amt = min(left, b["due"])
            part += 1
            conn.insert("payments", {"invoice_id": b["id"], "branch_id": branch_of(r["clinic"]) or b["branch_id"], "kind": "payment",
                                     "amount_cents": amt, "method": method, "account": account, "reference": r["ref"][:60],
                                     "received_at": f"{r['date']} 12:00:00", "received_by": user_id, "status": "valid",
                                     "notes": note[:500], "created_at": ts, "legacy_key": r["key"] if part == 1 else f"{r['key']}.{part}"})
            b["due"] -= amt
            left -= amt
            st["payments"] += 1
            st["paid_total"] += amt
            if dentist_id:
                for ln in b["lines"]:
                    if ln["dentist_id"] is None and same(ln["d"], wanted):
                        conn.execute("UPDATE invoice_items SET dentist_id = ?, legacy_dentist = 1 WHERE id = ?", (dentist_id, ln["id"]))
                        ln["dentist_id"] = dentist_id
                        st["lines_dentist"] += 1
        if left > 0:      # paid more than the MyMedsPH bills here show: keep it as account credit
            branch = branch_of(r["clinic"]) or (plist[0]["branch_id"] if plist else None) or settings.get(legacy_billing.BRANCH_KEY, conn)
            if not branch:
                st["skipped"] += 1
                continue
            conn.insert("patient_credits", {"patient_id": pid, "branch_id": branch, "kind": "deposit", "amount_cents": left, "method": method,
                                            "reference": r["ref"][:60], "entry_date": r["date"], "notes": (note + " (more than the bills)")[:500],
                                            "created_by": user_id, "created_at": ts, "legacy_key": r["key"] if part == 0 else f"{r['key']}.c"})
            st["credit"] += 1
            st["credit_total"] += left
    return st


def undo(conn) -> dict:
    """Remove the payments and credits brought in from collection reports (bills stay)."""
    n = conn.scalar("SELECT COUNT(*) FROM payments WHERE legacy_key LIKE 'mmc:%'") or 0
    c = conn.scalar("SELECT COUNT(*) FROM patient_credits WHERE legacy_key LIKE 'mmc:%'") or 0
    conn.execute("DELETE FROM payments WHERE legacy_key LIKE 'mmc:%'")
    conn.execute("DELETE FROM patient_credits WHERE legacy_key LIKE 'mmc:%'")
    settings.put(legacy_billing.COLLECTIONS_KEY, False, None, conn)
    return {"payments": n, "credits": c}
