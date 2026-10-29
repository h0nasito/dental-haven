"""MyMedsPH billing reports (Excel) -> dentist, paid and balance of the MyMedsPH bills.

MyMedsPH Reports > Billing exports a workbook per dentist (any date range) with a "Billing Report" sheet: one row per
bill (Date, Patient, Dentist, Clinic/Lab, Remarks, Procedures, Gross, Discount, Net, Paid, Balance). Unlike the data
export it says who the dentist was and how much of each bill is paid, so it's the most reliable source for both.

Matching: rows of one patient on one date are one bill. The patient is found by name (the report has no patient ID);
when a name fits several patients, the one with a MyMedsPH bill that day (and the same amount) is used. Each bill is
matched to the converted MyMedsPH bills of that patient on that date (MM-...).

What confirming does (the clinic's choices):
- dentist: set on the bill's lines that have none (lines naming another dentist are listed, not changed);
- payments: payments already here with a real date (collection reports, cashier) are kept; "Paid in MyMedsPH"
  payments made from a status number are replaced; only the missing amount is added, dated on the bill date, so the
  bill's paid and balance match MyMedsPH. Bills paid more here than in MyMedsPH are listed, not changed;
- bills in the report that aren't here are created (MB-...), with the report's procedures, net amount, dentist and paid.
Uploading the same or an overlapping report again changes nothing twice. Undo removes what the import added.
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import defaultdict

from . import legacy_billing, settings
from .billing import LEGACY_METHOD
from .legacy_collections import CollectionFileError, _branch_for, _cells, _date, _num, norm, patient_index
from .util import now_str

REPORT_SHEET = "Billing Report"
HEAD = ["date", "patient", "dentist", "clinic/lab", "remarks", "procedures", "gross", "discount", "net", "paid", "balance"]
CREATED_NOTE = "Imported from MyMedsPH billing report"


def read_reports(files: list[tuple[str, bytes]]) -> dict:
    """Read MyMedsPH billing report workbooks -> {"rows": [...], "files": [...], "errors": [...]}."""
    try:
        import openpyxl
    except ImportError as exc:
        raise CollectionFileError("The server can't read Excel files yet. Run the update on the server (dental-haven-update), "
                                  "then try again.") from exc
    rows, info, errors = [], [], []
    seen: dict[str, int] = defaultdict(int)
    for name, data in files:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable file
            raise CollectionFileError(f"{name}: this isn't an Excel file MyMedsPH made ({exc.__class__.__name__}).") from exc
        if REPORT_SHEET not in wb.sheetnames:
            raise CollectionFileError(f"{name}: no '{REPORT_SHEET}' sheet. Export it from MyMedsPH Reports > Billing.")
        it = wb[REPORT_SHEET].iter_rows(values_only=True)
        head = [c.lower() for c in _cells(next(it, []))]
        if head[:11] != HEAD:
            raise CollectionFileError(f"{name}: the '{REPORT_SHEET}' sheet doesn't have the usual columns ({', '.join(HEAD)}).")
        n = net_total = 0
        for line, r in enumerate(it, start=2):
            c = (_cells(r) + [""] * 11)[:11]
            if not any(c):
                continue
            d = _date(c[0])
            gross, disc, net, paid, bal = (_num(x) if x else 0 for x in c[6:11])
            if not d or not c[1] or net is None or paid is None:
                errors.append(f"{name} line {line}: missing date, patient or amounts - skipped")
                continue
            base = "|".join([d.isoformat(), norm(c[1]), norm(c[2]), norm(c[5]), str(net)])
            seen[base] += 1
            rows.append({"date": d.isoformat(), "patient": c[1], "dentist": c[2], "clinic": c[3], "procedures": c[5],
                         "gross": gross or 0, "discount": disc or 0, "net": net, "paid": paid, "balance": bal or 0,
                         "key": "mmb:" + hashlib.sha1(f"{base}#{seen[base]}".encode()).hexdigest()[:24]})
            n += 1
            net_total += net
        info.append({"file": name, "rows": n, "net": net_total})
    return {"rows": rows, "files": info, "errors": errors}


def split_procedures(text: str) -> list[tuple[str, str]]:
    """'Braces Adjustment(#), Extraction (#14, 15)' -> [('Braces Adjustment', ''), ('Extraction', '14, 15')]."""
    out = []
    for part in re.split(r",\s*(?![^()]*\))", text or ""):
        part = part.strip()
        if not part:
            continue
        m = re.search(r"\(#([^)]*)\)", part)
        tooth = (m.group(1).strip() if m else "")
        desc = re.sub(r"\s*\(#[^)]*\)", "", part).strip(" -") or part
        out.append((desc[:200], tooth[:60]))
    return out


def _bill_groups(rows) -> dict[tuple, dict]:
    """Rows of one patient on one date -> one bill."""
    groups: dict[tuple, dict] = {}
    for r in rows:
        g = groups.setdefault((norm(r["patient"]), r["date"]), {"patient": r["patient"], "date": r["date"], "clinic": r["clinic"],
                                                                 "net": 0, "paid": 0, "gross": 0, "discount": 0, "rows": [], "keys": []})
        for k in ("net", "paid", "gross", "discount"):
            g[k] += r[k]
        g["rows"].append(r)
        g["keys"].append(r["key"])
    return groups


def _is_synthetic(p) -> bool:
    """A 'Paid in MyMedsPH' payment made from a status number (no real date or amount behind it)."""
    return p["method"] == LEGACY_METHOD and not p["legacy_key"]


def plan(conn, data: dict) -> dict:
    """Work out what confirming would do, without changing anything."""
    idx = patient_index(conn)
    users = {norm(u["name"]): u["id"] for u in conn.all("SELECT id, name FROM users WHERE role = 'dentist'")}
    # MyMedsPH bills here, by patient and day (converted ones, and lines not converted yet)
    legacy: dict[tuple, list] = defaultdict(list)
    for b in conn.all("SELECT id, patient_id, branch_id, substr(issued_at, 1, 10) AS day, total_cents FROM invoices "
                      "WHERE legacy = 1 AND status = 'issued'"):
        legacy[(b["patient_id"], b["day"])].append(dict(b))
    pending = defaultdict(int)
    for r in conn.all("SELECT patient_id, bill_date AS day, SUM(total_cents) AS t FROM legacy_bills WHERE invoice_id IS NULL GROUP BY patient_id, bill_date"):
        pending[(r["patient_id"], r["day"])] += r["t"] or 0
    newbills = {(r["patient_id"], r["day"]) for r in conn.all(
        "SELECT patient_id, substr(issued_at, 1, 10) AS day FROM invoices WHERE legacy = 0 AND status = 'issued'")}
    pays: dict[int, list] = defaultdict(list)
    for p in conn.all("SELECT id, invoice_id, kind, amount_cents, method, legacy_key FROM payments WHERE status = 'valid' AND invoice_id IN "
                      "(SELECT id FROM invoices WHERE legacy = 1)"):
        pays[p["invoice_id"]].append(dict(p))
    lines: dict[int, list] = defaultdict(list)
    for ln in conn.all("SELECT id, invoice_id, description, dentist_id FROM invoice_items WHERE invoice_id IN (SELECT id FROM invoices WHERE legacy = 1)"):
        lines[ln["invoice_id"]].append(dict(ln))

    out = {"bills": [], "unmatched": defaultdict(lambda: [0, 0, ""]), "dentists_unmatched": set()}
    for g in _bill_groups(data["rows"]).values():
        ids = idx.get(norm(g["patient"]), [])
        day = g["date"]

        def here(pid):
            return legacy.get((pid, day), [])

        with_bill = [p for p in ids if here(p) or pending.get((p, day))]
        if len(with_bill) > 1:
            same = [p for p in with_bill if sum(b["total_cents"] for b in here(p)) + pending.get((p, day), 0) == g["net"]]
            with_bill = same if len(same) == 1 else with_bill
        if len(with_bill) == 1:
            pid, action = with_bill[0], "match"
        elif not with_bill and len(ids) == 1:
            pid, action = ids[0], ("billed_here" if (ids[0], day) in newbills else "create")
        else:
            u = out["unmatched"][(g["patient"], "several patients" if ids else "no patient")]
            u[0] += 1
            u[1] += g["net"]
            continue
        dent = {}
        for r in g["rows"]:
            if r["dentist"]:
                did = users.get(norm(r["dentist"]))
                if did:
                    dent[r["key"]] = did
                else:
                    out["dentists_unmatched"].add(r["dentist"])
        bill = {**g, "patient_id": pid, "action": action, "dentists": dent, "invoices": [], "ours_total": 0, "ours_paid": 0,
                "real_paid": 0, "synthetic": [], "add": 0, "more_here": 0, "dentist_lines": [], "dentist_conflicts": 0,
                "pending_total": pending.get((pid, day), 0)}
        if action == "match":
            invs = here(pid)
            bill["invoices"] = invs
            bill["ours_total"] = sum(b["total_cents"] for b in invs) + bill["pending_total"]
            for b in invs:
                for p in pays.get(b["id"], []):
                    amt = p["amount_cents"] if p["kind"] == "payment" else -p["amount_cents"]
                    bill["ours_paid"] += amt
                    if _is_synthetic(p):
                        bill["synthetic"].append(p)
                    else:
                        bill["real_paid"] += amt
            diff = g["paid"] - bill["real_paid"]
            bill["add"] = max(0, diff)
            bill["more_here"] = max(0, -diff)
            # dentist on lines with none: one dentist for the bill -> every empty line; else by procedure name
            names = {did for did in dent.values()}
            wanted = [(norm(d), dent.get(r["key"])) for r in g["rows"] for d, _t in split_procedures(r["procedures"])]
            for b in invs:
                for ln in lines.get(b["id"], []):
                    if len(names) == 1:
                        target = next(iter(names))
                    else:
                        dn = norm(ln["description"])
                        target = next((did for w, did in wanted if did and w and (dn.startswith(w) or w.startswith(dn))), None)
                    if not target:
                        continue
                    if ln["dentist_id"] is None:
                        bill["dentist_lines"].append((ln["id"], target))
                    elif ln["dentist_id"] != target:
                        bill["dentist_conflicts"] += 1
        out["bills"].append(bill)
    return out


def preview(conn, data: dict) -> dict:
    p = plan(conn, data)
    bills = p["bills"]
    rows = data["rows"]
    m = [b for b in bills if b["action"] == "match"]
    c = [b for b in bills if b["action"] == "create"]
    s = {"files": data["files"], "errors": data["errors"], "rows": len(rows), "bills": len(_bill_groups(rows)),
         "gross": sum(r["gross"] for r in rows), "discount": sum(r["discount"] for r in rows), "net": sum(r["net"] for r in rows),
         "paid": sum(r["paid"] for r in rows), "balance": sum(r["balance"] for r in rows),
         "matched": len(m), "matched_net": sum(b["net"] for b in m),
         "same_amount": sum(1 for b in m if b["ours_total"] == b["net"]),
         "amount_differs": sum(1 for b in m if b["ours_total"] != b["net"]),
         "create": len(c), "create_net": sum(b["net"] for b in c), "create_paid": sum(b["paid"] for b in c),
         "add": sum(1 for b in m if b["add"]), "add_total": sum(b["add"] for b in m),
         "synthetic": sum(len(b["synthetic"]) for b in m), "synthetic_total": sum(x["amount_cents"] for b in m for x in b["synthetic"]),
         "more_here": sum(1 for b in m if b["more_here"]), "more_here_total": sum(b["more_here"] for b in m),
         "dentist_lines": sum(len(b["dentist_lines"]) for b in m), "dentist_conflicts": sum(b["dentist_conflicts"] for b in m),
         "billed_here": sum(1 for b in bills if b["action"] == "billed_here"),
         "unmatched_names": sorted(((k[0], k[1], v[0], v[1]) for k, v in p["unmatched"].items()), key=lambda x: -x[3]),
         "dentists_unmatched": sorted(p["dentists_unmatched"])}
    s["unmatched_bills"] = sum(x[2] for x in s["unmatched_names"])
    s["unmatched_net"] = sum(x[3] for x in s["unmatched_names"])
    return s


def problems(conn, data: dict) -> list[list]:
    """Rows for the 'check these' download: bills whose amount or payments differ, and names not matched."""
    out = []
    for b in plan(conn, data)["bills"]:
        if b["action"] == "match" and (b["ours_total"] != b["net"] or b["more_here"] or b["dentist_conflicts"]):
            why = []
            if b["ours_total"] != b["net"]:
                why.append("amount differs")
            if b["more_here"]:
                why.append("paid more here than in MyMedsPH")
            if b["dentist_conflicts"]:
                why.append("another dentist on a line")
            out.append([b["patient"], b["date"], "; ".join(why), b["net"] / 100, b["ours_total"] / 100, b["paid"] / 100, b["ours_paid"] / 100])
        elif b["action"] == "billed_here":
            out.append([b["patient"], b["date"], "not in the MyMedsPH bills here; this day already has a bill made in this system",
                        b["net"] / 100, "", b["paid"] / 100, ""])
    for name, why, n, amt in preview(conn, data)["unmatched_names"]:
        out.append([name, "", f"name not matched ({why}), {n} bill(s)", amt / 100, "", "", ""])
    return out


def apply(conn, data: dict, user_id: int) -> dict:
    """Do it. Call inside a transaction."""
    st = {"converted": 0, "dentist_lines": 0, "payments": 0, "paid_total": 0, "replaced": 0, "created": 0, "created_net": 0,
          "skipped": 0}
    # MyMedsPH lines not turned into bills yet become bills first (unpaid; the report says what's paid)
    smap = legacy_billing.saved_map(conn)
    codes = [r["c"] for r in conn.all("SELECT DISTINCT status_code AS c FROM legacy_bills WHERE invoice_id IS NULL")]
    mapping = {c: ("history" if smap.get(c) == "history" else "unpaid") for c in codes}
    if mapping:
        st["converted"] = legacy_billing.convert(conn, mapping, settings.get(legacy_billing.BRANCH_KEY, conn), user_id, with_payments=False)["bills"]
        settings.put(legacy_billing.MAP_KEY, {**smap, **mapping}, user_id, conn)
    p = plan(conn, data)
    ts = now_str()
    branch_of = _branch_for(conn)
    have = {r["legacy_key"] for r in conn.all("SELECT legacy_key FROM payments WHERE legacy_key LIKE 'mmb:%'")}
    seq = conn.scalar("SELECT COUNT(*) FROM invoices WHERE number LIKE 'MB-%'") or 0
    for b in p["bills"]:
        if b["action"] == "match":
            for line_id, did in b["dentist_lines"]:
                conn.execute("UPDATE invoice_items SET dentist_id = ?, legacy_dentist = 2 WHERE id = ? AND dentist_id IS NULL", (did, line_id))
                st["dentist_lines"] += 1
            # status-made "Paid in MyMedsPH" payments give way to the report (kept as void, so Undo can bring them back)
            for x in b["synthetic"]:
                conn.execute("UPDATE payments SET status = 'void', legacy_key = ?, notes = ? WHERE id = ?",
                             (f"mmbv:{x['id']}", "Replaced by the MyMedsPH billing report", x["id"]))
                st["replaced"] += 1
            left = b["add"]
            if left <= 0:
                continue
            if not b["invoices"]:
                st["skipped"] += 1          # the MyMedsPH lines couldn't become a bill (patient has no branch)
                continue
            key = b["keys"][0]
            for inv in sorted(b["invoices"], key=lambda i: i["id"]):
                paid_here = conn.scalar("SELECT COALESCE(SUM(CASE WHEN kind = 'payment' THEN amount_cents ELSE -amount_cents END), 0) "
                                        "FROM payments WHERE invoice_id = ? AND status = 'valid'", (inv["id"],)) or 0
                due = inv["total_cents"] - paid_here
                amt = min(left, due) if due > 0 else 0
                if amt <= 0:
                    continue
                k = key if key not in have else f"{key}.{inv['id']}"
                if k in have:
                    continue
                conn.insert("payments", {"invoice_id": inv["id"], "branch_id": inv["branch_id"], "kind": "payment", "amount_cents": amt,
                                         "method": LEGACY_METHOD, "account": "", "reference": "", "received_at": f"{b['date']} 12:00:00",
                                         "received_by": user_id, "status": "valid", "notes": "Paid in MyMedsPH (billing report)",
                                         "created_at": ts, "legacy_key": k})
                have.add(k)
                left -= amt
                st["payments"] += 1
                st["paid_total"] += amt
        elif b["action"] == "create":
            pat = conn.one("SELECT preferred_branch_id FROM patients WHERE id = ?", (b["patient_id"],))
            branch = branch_of(b["clinic"]) or (pat["preferred_branch_id"] if pat else None) or settings.get(legacy_billing.BRANCH_KEY, conn)
            if not branch or b["net"] < 0:
                st["skipped"] += 1
                continue
            seq += 1
            inv = conn.insert("invoices", {"number": f"MB-{seq:06d}", "branch_id": branch, "patient_id": b["patient_id"], "status": "issued",
                                           "issued_at": f"{b['date']} 12:00:00", "subtotal_cents": b["net"], "discount_cents": 0, "tax_cents": 0,
                                           "total_cents": b["net"], "notes": CREATED_NOTE, "created_by": user_id, "created_at": ts, "legacy": 1})
            for r in b["rows"]:
                procs = split_procedures(r["procedures"]) or [("Services", "")]
                did = b["dentists"].get(r["key"])
                share = [r["net"] // len(procs)] * len(procs)
                share[0] += r["net"] - sum(share)          # the report has one amount per bill: shared over its procedures
                for (desc, tooth), amt in zip(procs, share):
                    conn.insert("invoice_items", {"invoice_id": inv, "description": desc, "qty": 1, "unit_price_cents": amt, "discount_cents": 0,
                                                  "amount_cents": amt, "tooth": tooth, "dentist_id": did, "legacy_dentist": 2 if did else 0,
                                                  "done_on": b["date"]})
            if b["paid"] > 0:
                conn.insert("payments", {"invoice_id": inv, "branch_id": branch, "kind": "payment", "amount_cents": min(b["paid"], b["net"]),
                                         "method": LEGACY_METHOD, "account": "", "reference": "", "received_at": f"{b['date']} 12:00:00",
                                         "received_by": user_id, "status": "valid", "notes": "Paid in MyMedsPH (billing report)",
                                         "created_at": ts, "legacy_key": b["keys"][0]})
                st["payments"] += 1
                st["paid_total"] += min(b["paid"], b["net"])
            st["created"] += 1
            st["created_net"] += b["net"]
    st["visits_dentist"] = link_procedure_dentists(conn)
    settings.put("legacy.billing_report_imported", True, user_id, conn)
    return st


def link_procedure_dentists(conn) -> int:
    """Imported MyMedsPH visits (progress notes) have no dentist: take it from that day's bill lines. One dentist on the
    day's lines -> every visit line that day; several -> by procedure name. Returns how many visit lines were set."""
    by_day: dict[tuple, list] = defaultdict(list)
    for r in conn.all("SELECT i.patient_id, COALESCE(ii.done_on, substr(i.issued_at, 1, 10)) AS day, ii.description, ii.dentist_id "
                      "FROM invoice_items ii JOIN invoices i ON i.id = ii.invoice_id WHERE i.status = 'issued' AND ii.dentist_id IS NOT NULL"):
        by_day[(r["patient_id"], r["day"])].append((norm((r["description"] or "").split(" · ")[0]), r["dentist_id"]))
    n = 0
    for p in conn.all("SELECT id, patient_id, substr(performed_at, 1, 10) AS day, description FROM procedures "
                      "WHERE legacy_key IS NOT NULL AND performed_by IS NULL AND status = 'completed'"):
        lines = by_day.get((p["patient_id"], p["day"]))
        if not lines:
            continue
        dentists = {d for _, d in lines}
        if len(dentists) == 1:
            target = next(iter(dentists))
        else:
            name = norm((p["description"] or "").split("\n", 1)[0])
            target = next((d for w, d in lines if w and name and (w.startswith(name) or name.startswith(w))), None)
        if target:
            conn.execute("UPDATE procedures SET performed_by = ?, legacy_dentist = 1 WHERE id = ?", (target, p["id"]))
            n += 1
    return n


def status(conn) -> dict:
    return {"payments": conn.scalar("SELECT COUNT(*) FROM payments WHERE legacy_key LIKE 'mmb:%'") or 0,
            "paid_total": conn.scalar("SELECT COALESCE(SUM(amount_cents), 0) FROM payments WHERE legacy_key LIKE 'mmb:%'") or 0,
            "created": conn.scalar("SELECT COUNT(*) FROM invoices WHERE notes = ?", (CREATED_NOTE,)) or 0,
            "lines": conn.scalar("SELECT COUNT(*) FROM invoice_items WHERE legacy_dentist = 2") or 0}


def undo(conn) -> dict:
    """Remove what the billing report import added: its payments, the bills it created (if nobody worked on them), the
    dentists it set; and bring back the status-made payments it replaced."""
    st = status(conn)
    keep = {r["invoice_id"] for r in conn.all("SELECT invoice_id FROM payments WHERE status = 'valid' AND (legacy_key IS NULL OR legacy_key NOT LIKE 'mmb:%') "
                                               "AND invoice_id IN (SELECT id FROM invoices WHERE notes = ?)", (CREATED_NOTE,))}
    conn.execute("DELETE FROM payments WHERE legacy_key LIKE 'mmb:%'")
    created = [r["id"] for r in conn.all("SELECT id FROM invoices WHERE notes = ?", (CREATED_NOTE,)) if r["id"] not in keep]
    for iid in created:
        conn.execute("DELETE FROM invoice_items WHERE invoice_id = ?", (iid,))
        conn.execute("DELETE FROM invoices WHERE id = ?", (iid,))
    conn.execute("UPDATE invoice_items SET dentist_id = NULL, legacy_dentist = 0 WHERE legacy_dentist = 2")
    conn.execute("UPDATE procedures SET performed_by = NULL, legacy_dentist = 0 WHERE legacy_dentist = 1")
    restored = conn.scalar("SELECT COUNT(*) FROM payments WHERE legacy_key LIKE 'mmbv:%'") or 0
    conn.execute("UPDATE payments SET status = 'valid', legacy_key = NULL, notes = 'Paid in MyMedsPH (imported)' WHERE legacy_key LIKE 'mmbv:%'")
    settings.put("legacy.billing_report_imported", False, None, conn)
    return {**st, "removed_bills": len(created), "kept_bills": len(keep), "restored": restored}
