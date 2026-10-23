"""MyMedsPH collection reports -> payments on the converted MyMedsPH bills (synthetic data only)."""
from __future__ import annotations

import io
import secrets
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_app import Base  # noqa: E402

HEAD = ["Date", "Patient", "Clinic/Lab", "Reference #", "Remarks", "Amount"]


def workbook(rows, dentist_rows=None):
    import openpyxl
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Collection Summary"
    s.append(["Associates", "Clinic", "Laboratory", "Payment Type", "Date From", "Date To", "Total Amount", "Total Row"])
    r = wb.create_sheet("Collection Report")
    r.append(HEAD)
    for row in rows:
        r.append(row)
    for name, drows in (dentist_rows or {}).items():
        d = wb.create_sheet(name)
        d.append(HEAD + ["Procedures"])
        for row in drows:
            d.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestLegacyCollections(Base):
    def setUp(self):
        from app.util import now_str
        for t in ("payments", "patient_credits"):
            self.conn.execute(f"DELETE FROM {t} WHERE legacy_key LIKE 'mmc:%'")
        self.conn.execute("DELETE FROM payments WHERE invoice_id IN (SELECT id FROM invoices WHERE legacy = 1)")
        self.conn.execute("DELETE FROM invoice_items WHERE invoice_id IN (SELECT id FROM invoices WHERE legacy = 1)")
        self.conn.execute("UPDATE legacy_bills SET invoice_id = NULL")
        self.conn.execute("DELETE FROM invoices WHERE legacy = 1")
        self.conn.execute("DELETE FROM legacy_bills")
        self.conn.execute("DELETE FROM settings WHERE key LIKE 'legacy.%'")
        tag = secrets.token_hex(3)
        self.first = f"Paya{tag}"
        b = self.branch("malolos")
        self.pid = self.conn.insert("patients", {"first_name": f"{self.first} rose", "middle_name": "Mid", "last_name": "Collqt",
                                                 "chart_no": f"DH-C{tag}", "preferred_branch_id": b, "created_at": now_str(), "updated_at": now_str()})
        # two patients with the same name -> ambiguous
        for i in (1, 2):
            self.conn.insert("patients", {"first_name": f"Twin{tag}", "last_name": "Same", "chart_no": f"DH-T{i}{tag}", "preferred_branch_id": b,
                                          "created_at": now_str(), "updated_at": now_str()})
        self.twin = f"Twin{tag} Same"
        bills = [("2026-10-01", "PANORAMIC XRAY", 0), ("2026-10-01", "CBCT SCAN", 850000), ("2026-10-01", "Denture Impression", 3500000),
                 ("2026-10-09", "IMPLANT", 40000000), ("2026-10-09", "BONE GRAFT - Per Unit", 15500000), ("2026-10-13", "Extraction", 300000)]
        for i, (d, svc, tot) in enumerate(bills):
            self.conn.insert("legacy_bills", {"patient_id": self.pid, "bill_date": d, "service": svc, "qty": 1, "unit_price_cents": tot,
                                              "total_cents": tot, "status_code": "2", "legacy_key": f"cq-{tag}-{i}"})
        self.name = f"{self.first} rose Collqt"

    def upload(self, c, data):
        return c.post("/staff/admin/import/collections", data={"files": [(io.BytesIO(data), "2026.xlsx")]}, content_type="multipart/form-data")

    def test_payments_match_mymeds_balance(self):
        rows = [["Oct 01, 2026", self.name, "Dental Haven Malolos", "cash", "-", "8500.0"],
                ["Oct 01, 2026", self.name, "Dental Haven Malolos", "BDO", "-downpayment", "17500.0"],
                ["Oct 09, 2026", self.name, "Dental Haven Malolos", "gcash", "-", "385000.0"],
                ["Oct 13, 2026", self.name, "Dental Haven Malolos", "cash", "-", "3000.0"],
                ["Oct 14, 2026", "Nobody Here", "Dental Haven Malolos", "cash", "-", "500.0"],
                ["Oct 14, 2026", self.twin, "Dental Haven Malolos", "cash", "-", "700.0"]]
        dent = {"Dr. Lea Ramos (demo)": [rows[3] + ["Extraction- Anterior/Posterior(#13)"]]}
        c = self.login("admin")
        r = self.upload(c, workbook(rows, dent))
        page = r.data.decode()
        self.assertEqual(r.status_code, 200, page[:500])
        self.assertIn("Nobody Here", page)
        self.assertIn("several patients", page)
        token = page.split('name="token" value="')[1].split('"')[0]
        r = c.post("/staff/admin/import/collections/confirm", data={"token": token, "action": "unmatched"})
        self.assertIn("Nobody Here", r.data.decode("utf-8-sig"))
        r = c.post("/staff/admin/import/collections/confirm", data={"token": token, "action": "confirm"})
        self.assertEqual(r.status_code, 302)
        billed = self.q("SELECT SUM(total_cents) AS t FROM invoices WHERE patient_id = ? AND legacy = 1", (self.pid,))["t"]
        paid = self.q("SELECT SUM(p.amount_cents) AS t FROM payments p JOIN invoices i ON i.id = p.invoice_id WHERE i.patient_id = ?", (self.pid,))["t"]
        self.assertEqual((billed, paid), (60150000, 41400000))
        self.assertEqual(billed - paid, 18750000)                     # same as MyMedsPH: P187,500
        # the Oct 13 extraction payment went to the Oct 13 bill, and the dentist was set on that line
        ext = self.q("SELECT ii.dentist_id, i.id FROM invoice_items ii JOIN invoices i ON i.id = ii.invoice_id WHERE i.patient_id = ? AND ii.description = 'Extraction'", (self.pid,))
        self.assertIsNotNone(ext["dentist_id"])
        self.assertEqual(self.q("SELECT SUM(amount_cents) AS t FROM payments WHERE invoice_id = ?", (ext["id"],))["t"], 300000)
        pay = self.q("SELECT * FROM payments WHERE legacy_key LIKE 'mmc:%' AND method = 'bank_transfer'")
        self.assertEqual(pay["account"], "BDO")
        # the same file again adds nothing
        r = self.upload(c, workbook(rows, dent))
        token = r.data.decode().split('name="token" value="')[1].split('"')[0]
        c.post("/staff/admin/import/collections/confirm", data={"token": token, "action": "confirm"})
        self.assertEqual(self.q("SELECT SUM(p.amount_cents) AS t FROM payments p JOIN invoices i ON i.id = p.invoice_id WHERE i.patient_id = ?",
                                (self.pid,))["t"], 41400000)
        # paying more than all bills -> account credit
        extra = rows + [["Oct 20, 2026", self.name, "Dental Haven Malolos", "cash", "-", "200000.0"]]
        r = self.upload(c, workbook(extra, dent))
        token = r.data.decode().split('name="token" value="')[1].split('"')[0]
        c.post("/staff/admin/import/collections/confirm", data={"token": token, "action": "confirm"})
        self.assertEqual(self.q("SELECT SUM(amount_cents) AS t FROM patient_credits WHERE patient_id = ? AND legacy_key LIKE 'mmc:%'", (self.pid,))["t"],
                         20000000 - 18750000)
        # undo the bills removes the bills and their report payments
        c.post("/staff/admin/import/bills", data={"action": "undo"})
        self.assertEqual(self.q("SELECT COUNT(*) AS n FROM invoices WHERE patient_id = ? AND legacy = 1", (self.pid,))["n"], 0)
        self.assertEqual(self.q("SELECT COUNT(*) AS n FROM payments WHERE legacy_key LIKE 'mmc:%'")["n"], 0)

    def test_not_a_report(self):
        c = self.login("admin")
        r = c.post("/staff/admin/import/collections", data={"files": [(io.BytesIO(b"hello"), "x.xlsx")]}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302)
