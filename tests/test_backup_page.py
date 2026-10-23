"""Administration > Backup & download: full zip with database, files, Excel copy and README; choices; super admin only."""
from __future__ import annotations

import io
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_app import Base  # noqa: E402


class TestBackupPage(Base):
    def make(self, c, parts):
        """Prepare a backup (runs at once in tests), then download it from the link on the page."""
        import re
        self.assertEqual(c.post("/staff/admin/backup", data={"parts": parts}).status_code, 302)
        page = c.get("/staff/admin/backup").data.decode()
        link = re.findall(r'href="(/staff/admin/backup/[0-9a-f]+/download)"', page)[0]
        return c.get(link)

    def test_full_backup_contents(self):
        c = self.login("admin")
        page = c.get("/staff/admin/backup").data.decode()
        for part in ("Backup &amp; download", "Database", "Uploaded files", "Excel copy", "Patients"):
            self.assertIn(part, page)
        r = self.make(c, ["database", "files", "excel"])
        self.assertEqual((r.status_code, r.mimetype), (200, "application/zip"))
        z = zipfile.ZipFile(io.BytesIO(r.data))
        names = z.namelist()
        for n in ("dental_haven.db", "dental-haven-data.xlsx", "README.txt"):
            self.assertIn(n, names)
        self.assertIn("How to restore", z.read("README.txt").decode())
        # the database copy opens and has the patients
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "db.sqlite"
            p.write_bytes(z.read("dental_haven.db"))
            db = sqlite3.connect(p)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM patients").fetchone()[0], self.q("SELECT COUNT(*) AS n FROM patients")["n"])
            db.close()
            from openpyxl import load_workbook
            x = Path(d) / "data.xlsx"
            x.write_bytes(z.read("dental-haven-data.xlsx"))
            wb = load_workbook(x, read_only=True)
            self.assertIn("Contents", wb.sheetnames)
            self.assertIn("patients", wb.sheetnames)
            self.assertNotIn("sessions", wb.sheetnames)
            users = next(wb["users"].iter_rows(max_row=1, values_only=True))
            self.assertIn("email", users)
            self.assertNotIn("password_hash", users)
            wb.close()
        self.assertIn("Downloaded a backup", self.q("SELECT summary FROM audit_log WHERE action = 'backup_downloaded' ORDER BY id DESC")["summary"])
        self.assertIn("Recent downloads", c.get("/staff/admin/backup").data.decode())
        # delete it from the server: the link stops working
        import re
        job = re.findall(r'name="job" value="([0-9a-f]+)"', c.get("/staff/admin/backup").data.decode())[0]
        c.post("/staff/admin/backup", data={"action": "delete", "job": job})
        self.assertEqual(c.get(f"/staff/admin/backup/{job}/download").status_code, 302)

    def test_only_excel_and_nothing_chosen(self):
        c = self.login("admin")
        names = zipfile.ZipFile(io.BytesIO(self.make(c, ["excel"]).data)).namelist()
        self.assertEqual(sorted(names), ["README.txt", "dental-haven-data.xlsx"])
        self.assertEqual(c.post("/staff/admin/backup", data={}).status_code, 302)

    def test_staff_cannot(self):
        s = self.login("staff.malolos")
        self.assertEqual(s.get("/staff/admin/backup").status_code, 403)
        self.assertEqual(s.post("/staff/admin/backup", data={"parts": ["database"]}).status_code, 403)
        self.assertEqual(s.get("/staff/admin/backup/abc/download").status_code, 403)
